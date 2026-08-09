"""Build the Excel review workbook from the generated aggregates.

Produces ``05_dashboard/san_diego_ops_review.xlsx`` with six sheets:

    Executive Summary   headline metrics with live formulas
    Monthly KPI Review  month-by-month demand and closure activity
    Service Categories  full active-backlog breakdown, filterable
    Geography           council district and community detail
    Aging Analysis      bucket profile and the category-by-district hotspots
    Data Quality        every audit check with its grade

Honest statement about PivotTables: openpyxl cannot author a native Excel
PivotTable cache from scratch, so this workbook contains **no native
PivotTables**. It uses Excel Tables (structured references, filter dropdowns,
banded rows) plus native charts, conditional formatting and live formulas. Any
sheet can be turned into a PivotTable by the reader in two clicks because the
data is already in tabular form. This is stated in the workbook itself, on the
Executive Summary sheet.

Usage
-----
    python -m src.build_excel
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    from src import claims as claims_mod, config
except ImportError:  # pragma: no cover
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from src import claims as claims_mod, config

# --- shared styling ----------------------------------------------------------
INK = "FF0B0B0B"
MUTED = "FF52514E"
ACCENT = "FF2A78D6"
ACCENT_2 = "FFEB6834"
HEADER_FILL = PatternFill("solid", fgColor="FF1F3B5C")
TITLE_FONT = Font(size=15, bold=True, color=INK)
SUB_FONT = Font(size=10, color=MUTED)
H2_FONT = Font(size=11, bold=True, color=INK)
HEADER_FONT = Font(size=10, bold=True, color="FFFFFFFF")
NOTE_FONT = Font(size=9, color=MUTED, italic=True)
THIN = Side(style="thin", color="FFD9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BOUNDARY_NOTE = (
    "Data boundary: Get It Done records represent submitted service requests and case "
    "statuses, not verified maintenance completion. A closed case records a case closure, "
    "not a completed repair. Nothing in this workbook measures crew performance, and no "
    "relationship shown is causal."
)


def write_table(ws, df: pd.DataFrame, start_row: int, name: str,
                number_formats: dict[str, str] | None = None) -> tuple[int, int]:
    """Write a DataFrame as a native Excel Table. Returns (first_row, last_row)."""
    number_formats = number_formats or {}
    header_row = start_row

    for j, column in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=j, value=str(column))
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX

    for i, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        for j, value in enumerate(row, start=1):
            if isinstance(value, (pd.Timestamp, datetime)):
                value = value.strftime("%Y-%m-%d")
            elif pd.isna(value):
                value = None
            elif hasattr(value, "item"):
                value = value.item()
            cell = ws.cell(row=i, column=j, value=value)
            cell.border = BOX
            fmt = number_formats.get(df.columns[j - 1])
            if fmt:
                cell.number_format = fmt

    last_row = header_row + len(df)
    ref = f"A{header_row}:{get_column_letter(len(df.columns))}{last_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight9", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)

    for j, column in enumerate(df.columns, start=1):
        longest = max([len(str(column))] + [len(str(v)) for v in df.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, 11), 46)
    ws.row_dimensions[header_row].height = 30
    return header_row, last_row


def title_block(ws, title: str, subtitle: str, question: str = "") -> int:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUB_FONT
    row = 3
    if question:
        ws[f"A{row}"] = question
        ws[f"A{row}"].font = NOTE_FONT
        row += 1
    return row + 1


def build() -> Path:
    c = claims_mod.build_claims()
    agg = claims_mod.agg
    meta = json.loads(config.RUN_METADATA.read_text())
    snapshot = meta["snapshot_date"]
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    wb = Workbook()

    # =====================================================================
    # 1. Executive Summary
    # =====================================================================
    ws = wb.active
    ws.title = "Executive Summary"
    ws["A1"] = "San Diego Service Operations Intelligence"
    ws["A1"].font = Font(size=16, bold=True, color=INK)
    ws["A2"] = f"Active backlog review — data snapshot {snapshot}"
    ws["A2"].font = SUB_FONT
    ws["A3"] = f"Generated {generated} from City of San Diego Open Data Portal extracts"
    ws["A3"].font = SUB_FONT

    ws["A5"] = BOUNDARY_NOTE
    ws["A5"].font = Font(size=10, bold=True, color="FF8A3B12")
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A5:G7")
    ws["A5"].fill = PatternFill("solid", fgColor="FFFDF0E8")

    ws["A9"] = "Headline metrics"
    ws["A9"].font = H2_FONT

    kpis = agg("agg_executive_kpis")
    kpis_display = kpis[kpis["metric_key"] != "snapshot_date"][
        ["metric_label", "value", "unit", "grouping"]
    ].rename(columns={
        "metric_label": "Metric", "value": "Value", "unit": "Unit", "grouping": "Group"})
    first, last = write_table(ws, kpis_display, 10, "tblExecKPIs",
                              {"Value": "#,##0.0"})

    ws.conditional_formatting.add(
        f"B{first + 1}:B{last}",
        DataBarRule(start_type="min", end_type="max", color="FF2A78D6", showValue=True))

    # Live cross-checks: the workbook re-derives three figures with formulas so a
    # reader can see the arithmetic rather than trust a pasted number.
    check_row = last + 3
    ws[f"A{check_row}"] = "Live consistency checks (Excel formulas, recalculated on open)"
    ws[f"A{check_row}"].font = H2_FONT

    checks = [
        ("Aged 90+ as % of active backlog",
         f'=ROUND(100*INDEX(tblExecKPIs[Value],MATCH("Active requests aged 90+ days",'
         f'tblExecKPIs[Metric],0))/INDEX(tblExecKPIs[Value],'
         f'MATCH("Active request backlog (case records)",tblExecKPIs[Metric],0)),1)',
         c["active_aged_90_plus_pct"].value),
        ("Duplicate children as % of active backlog",
         f'=ROUND(100*INDEX(tblExecKPIs[Value],MATCH("Active case records flagged as duplicate children",'
         f'tblExecKPIs[Metric],0))/INDEX(tblExecKPIs[Value],'
         f'MATCH("Active request backlog (case records)",tblExecKPIs[Metric],0)),1)',
         c["duplicate_rate_active_pct"].value),
        ("Year-over-year demand change %",
         f'=ROUND(100*(INDEX(tblExecKPIs[Value],MATCH("Submissions, January to last complete month, current year",'
         f'tblExecKPIs[Metric],0))-INDEX(tblExecKPIs[Value],'
         f'MATCH("Submissions, same months, prior year",tblExecKPIs[Metric],0)))/'
         f'INDEX(tblExecKPIs[Value],MATCH("Submissions, same months, prior year",'
         f'tblExecKPIs[Metric],0)),1)',
         c["demand_yoy_change_pct"].value),
    ]
    hdr = check_row + 1
    for j, label in enumerate(["Check", "Formula result", "Expected (from SQL)", "Status"], start=1):
        cell = ws.cell(row=hdr, column=j, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BOX
    for k, (label, formula, expected) in enumerate(checks, start=1):
        r = hdr + k
        ws.cell(row=r, column=1, value=label).border = BOX
        cell = ws.cell(row=r, column=2, value=formula)
        cell.number_format = "0.0"
        cell.border = BOX
        ws.cell(row=r, column=3, value=round(float(expected), 1)).border = BOX
        ws.cell(row=r, column=3).number_format = "0.0"
        status = ws.cell(row=r, column=4,
                         value=f'=IF(ROUND(B{r},1)=ROUND(C{r},1),"MATCH","REVIEW")')
        status.border = BOX
    ws.conditional_formatting.add(
        f"D{hdr + 1}:D{hdr + len(checks)}",
        CellIsRule(operator="equal", formula=['"MATCH"'],
                   fill=PatternFill("solid", fgColor="FFDDF3DD")))
    ws.conditional_formatting.add(
        f"D{hdr + 1}:D{hdr + len(checks)}",
        CellIsRule(operator="equal", formula=['"REVIEW"'],
                   fill=PatternFill("solid", fgColor="FFFBE0E0")))

    note_row = hdr + len(checks) + 2
    ws[f"A{note_row}"] = (
        "Note on PivotTables: this workbook contains no native Excel PivotTables. "
        "They cannot be authored reliably from Python, so rather than claim otherwise "
        "every sheet is a native Excel Table with filter dropdowns — select any of them "
        "and use Insert > PivotTable to build one directly.")
    ws[f"A{note_row}"].font = NOTE_FONT
    ws[f"A{note_row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(f"A{note_row}:G{note_row + 2}")
    ws.column_dimensions["A"].width = 52

    # =====================================================================
    # 2. Monthly KPI Review
    # =====================================================================
    ws = wb.create_sheet("Monthly KPI Review")
    row = title_block(
        ws, "Monthly demand and recorded closures",
        f"Snapshot {snapshot}. Coverage flags matter: months before January 2025 are not "
        f"fully covered by the extracts in scope.",
        "Q11 - How has demand changed where the source data allows a fair comparison?")

    monthly = agg("agg_demand_monthly")[
        ["month_start", "submissions", "distinct_issues", "duplicate_children",
         "still_active", "pct_still_active", "coverage_status"]
    ].rename(columns={
        "month_start": "Month", "submissions": "Submissions",
        "distinct_issues": "Distinct issues", "duplicate_children": "Duplicate children",
        "still_active": "Still active", "pct_still_active": "% still active",
        "coverage_status": "Coverage"})
    monthly["Month"] = pd.to_datetime(monthly["Month"]).dt.strftime("%Y-%m")
    first, last = write_table(ws, monthly, row, "tblMonthly",
                              {"Submissions": "#,##0", "Distinct issues": "#,##0",
                               "Duplicate children": "#,##0", "Still active": "#,##0",
                               "% still active": "0.0"})
    ws.conditional_formatting.add(
        f"B{first + 1}:B{last}",
        ColorScaleRule(start_type="min", start_color="FFFFFFFF",
                       end_type="max", end_color="FF9EC5F4"))

    chart = LineChart()
    chart.title = "Submissions by month"
    chart.height, chart.width = 8, 22
    chart.y_axis.title = "Case records"
    data = Reference(ws, min_col=2, min_row=first, max_row=last)
    cats = Reference(ws, min_col=1, min_row=first + 1, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"I{first}")

    closures = agg("agg_closure_monthly")[
        ["month_start", "closures_recorded", "closed_status", "referred_status",
         "median_lifecycle_days", "p90_lifecycle_days", "coverage_status"]
    ].rename(columns={
        "month_start": "Month", "closures_recorded": "Closures recorded",
        "closed_status": "Closed", "referred_status": "Referred",
        "median_lifecycle_days": "Median lifecycle (days)",
        "p90_lifecycle_days": "P90 lifecycle (days)", "coverage_status": "Coverage"})
    closures["Month"] = pd.to_datetime(closures["Month"]).dt.strftime("%Y-%m")
    row2 = last + 22
    ws.cell(row=row2 - 1, column=1, value="Recorded case closures by month").font = H2_FONT
    write_table(ws, closures, row2, "tblClosures",
                {"Closures recorded": "#,##0", "Closed": "#,##0", "Referred": "#,##0",
                 "Median lifecycle (days)": "#,##0", "P90 lifecycle (days)": "#,##0"})

    # =====================================================================
    # 3. Service Categories
    # =====================================================================
    ws = wb.create_sheet("Service Categories")
    row = title_block(
        ws, "Active backlog by service category",
        "Every category in the active queue. Use the filter dropdowns to isolate a group.",
        "Q2, Q3, Q4 - where the workload sits, which categories are oldest, "
        "and which have the longest tail.")

    service = agg("agg_service_backlog")[
        ["service_name", "most_common_record_type", "active_records", "active_distinct_issues",
         "duplicate_rate_pct", "pct_of_active_backlog", "median_age_days", "p90_age_days",
         "mean_age_days", "aged_60_plus", "aged_90_plus", "aged_365_plus",
         "pct_aged_90_plus", "pct_of_aged_90_backlog", "rank_by_volume", "rank_by_median_age"]
    ].rename(columns={
        "service_name": "Service category", "most_common_record_type": "Most common record type",
        "active_records": "Active records", "active_distinct_issues": "Distinct issues",
        "duplicate_rate_pct": "Duplicate rate %", "pct_of_active_backlog": "% of backlog",
        "median_age_days": "Median age (days)", "p90_age_days": "P90 age (days)",
        "mean_age_days": "Mean age (days)", "aged_60_plus": "Aged 60+",
        "aged_90_plus": "Aged 90+", "aged_365_plus": "Aged 365+",
        "pct_aged_90_plus": "% of own queue aged 90+",
        "pct_of_aged_90_backlog": "% of city aged 90+",
        "rank_by_volume": "Rank by volume", "rank_by_median_age": "Rank by median age"})
    first, last = write_table(ws, service, row, "tblService",
                              {"Active records": "#,##0", "Distinct issues": "#,##0",
                               "Duplicate rate %": "0.0", "% of backlog": "0.00",
                               "Median age (days)": "#,##0", "P90 age (days)": "#,##0",
                               "Mean age (days)": "#,##0.0", "Aged 60+": "#,##0",
                               "Aged 90+": "#,##0", "Aged 365+": "#,##0",
                               "% of own queue aged 90+": "0.0",
                               "% of city aged 90+": "0.00"})
    ws.conditional_formatting.add(
        f"C{first + 1}:C{last}",
        DataBarRule(start_type="min", end_type="max", color="FF2A78D6", showValue=True))
    ws.conditional_formatting.add(
        f"M{first + 1}:M{last}",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFFFF",
                       end_type="num", end_value=100, end_color="FFEB6834"))

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Active records by service category (top of table)"
    chart.height, chart.width = 12, 20
    top_last = min(first + 12, last)
    chart.add_data(Reference(ws, min_col=3, min_row=first, max_row=top_last),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=top_last))
    ws.add_chart(chart, f"R{first}")

    # =====================================================================
    # 4. Geography
    # =====================================================================
    ws = wb.create_sheet("Geography")
    row = title_block(
        ws, "Active backlog by geography",
        "District counts are not comparable as service levels: this dataset carries no "
        "population, street-mileage or asset denominators.",
        "Q5, Q6 - where volume sits and whether any area's queue is unusually old.")

    district = agg("agg_geography_district")[
        ["council_district", "active_records", "pct_of_active_backlog",
         "submissions_last_12_months", "backlog_per_recent_submission", "median_age_days",
         "p90_age_days", "aged_90_plus", "pct_aged_90_plus", "pct_of_aged_90_backlog",
         "duplicate_rate_pct"]
    ].rename(columns={
        "council_district": "Council district", "active_records": "Active records",
        "pct_of_active_backlog": "% of backlog",
        "submissions_last_12_months": "Submissions, last 12 months",
        "backlog_per_recent_submission": "Backlog per recent submission",
        "median_age_days": "Median age (days)", "p90_age_days": "P90 age (days)",
        "aged_90_plus": "Aged 90+", "pct_aged_90_plus": "% of own queue aged 90+",
        "pct_of_aged_90_backlog": "% of city aged 90+", "duplicate_rate_pct": "Duplicate rate %"})
    first, last = write_table(ws, district, row, "tblDistrict",
                              {"Active records": "#,##0", "% of backlog": "0.0",
                               "Submissions, last 12 months": "#,##0",
                               "Backlog per recent submission": "0.000",
                               "Median age (days)": "#,##0", "P90 age (days)": "#,##0",
                               "Aged 90+": "#,##0", "% of own queue aged 90+": "0.0",
                               "% of city aged 90+": "0.0", "Duplicate rate %": "0.0"})
    ws.conditional_formatting.add(
        f"E{first + 1}:E{last}",
        ColorScaleRule(start_type="min", start_color="FFFFFFFF",
                       end_type="max", end_color="FFEB6834"))

    index_df = agg("agg_geography_aging_index")
    index_df = index_df[index_df["area_type"] == "Council district"][
        ["area", "active_records", "aged_90_plus", "median_age_days",
         "pct_of_backlog", "pct_of_aged_90", "aged_concentration_index", "rank_by_index"]
    ].rename(columns={
        "area": "Council district", "active_records": "Active records",
        "aged_90_plus": "Aged 90+", "median_age_days": "Median age (days)",
        "pct_of_backlog": "% of backlog", "pct_of_aged_90": "% of aged 90+",
        "aged_concentration_index": "Aged concentration index", "rank_by_index": "Rank"})
    row2 = last + 3
    ws.cell(row=row2 - 1, column=1,
            value="Aged-concentration index (1.000 = area holds exactly the aged share "
                  "its queue size implies)").font = H2_FONT
    f2, l2 = write_table(ws, index_df, row2, "tblAgingIndex",
                         {"Active records": "#,##0", "Aged 90+": "#,##0",
                          "Median age (days)": "#,##0", "% of backlog": "0.00",
                          "% of aged 90+": "0.00", "Aged concentration index": "0.000"})
    ws.conditional_formatting.add(
        f"G{f2 + 1}:G{l2}",
        CellIsRule(operator="greaterThan", formula=["1"],
                   fill=PatternFill("solid", fgColor="FFFDE7DD")))

    community = agg("agg_geography_community")[
        ["community", "example_council_district", "active_records", "pct_of_active_backlog",
         "median_age_days", "p90_age_days", "aged_90_plus", "pct_aged_90_plus"]
    ].rename(columns={
        "community": "Community", "example_council_district": "Council district (example)",
        "active_records": "Active records", "pct_of_active_backlog": "% of backlog",
        "median_age_days": "Median age (days)", "p90_age_days": "P90 age (days)",
        "aged_90_plus": "Aged 90+", "pct_aged_90_plus": "% of own queue aged 90+"})
    row3 = l2 + 3
    ws.cell(row=row3 - 1, column=1,
            value="Community planning areas with 100+ active records").font = H2_FONT
    write_table(ws, community, row3, "tblCommunity",
                {"Active records": "#,##0", "% of backlog": "0.00",
                 "Median age (days)": "#,##0", "P90 age (days)": "#,##0",
                 "Aged 90+": "#,##0", "% of own queue aged 90+": "0.0"})

    # =====================================================================
    # 5. Aging Analysis
    # =====================================================================
    ws = wb.create_sheet("Aging Analysis")
    row = title_block(
        ws, "Age profile of the active queue",
        "Active age = snapshot date minus request date. It measures how long a request has "
        "been open, not how long any work took.",
        "Q1, Q12 - how old the queue is and which cells to look at first.")

    buckets = agg("agg_backlog_aging_buckets")[
        ["age_bucket", "n_records", "n_distinct_issues", "median_age_days",
         "pct_of_active_backlog", "cumulative_records", "cumulative_pct"]
    ].rename(columns={
        "age_bucket": "Age bucket (days)", "n_records": "Active records",
        "n_distinct_issues": "Distinct issues", "median_age_days": "Median age (days)",
        "pct_of_active_backlog": "% of backlog", "cumulative_records": "Cumulative records",
        "cumulative_pct": "Cumulative %"})
    first, last = write_table(ws, buckets, row, "tblBuckets",
                              {"Active records": "#,##0", "Distinct issues": "#,##0",
                               "Median age (days)": "#,##0", "% of backlog": "0.0",
                               "Cumulative records": "#,##0", "Cumulative %": "0.0"})
    ws.conditional_formatting.add(
        f"B{first + 1}:B{last}",
        DataBarRule(start_type="min", end_type="max", color="FFEB6834", showValue=True))

    chart = BarChart()
    chart.title = "Active records by age bucket"
    chart.height, chart.width = 9, 18
    chart.add_data(Reference(ws, min_col=2, min_row=first, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=last))
    ws.add_chart(chart, f"J{first}")

    priority = agg("agg_priority_table")[
        ["investigation_rank", "service_name", "active_records", "median_age_days",
         "p90_age_days", "aged_90_plus", "pct_of_own_queue_aged_90",
         "pct_of_scored_aged_90_backlog", "priority_score", "why_flagged"]
    ].rename(columns={
        "investigation_rank": "Rank", "service_name": "Service category",
        "active_records": "Active records", "median_age_days": "Median age (days)",
        "p90_age_days": "P90 age (days)", "aged_90_plus": "Aged 90+",
        "pct_of_own_queue_aged_90": "% of own queue aged 90+",
        "pct_of_scored_aged_90_backlog": "% of city aged 90+",
        "priority_score": "Priority score", "why_flagged": "Why flagged"})
    row2 = last + 20
    ws.cell(row=row2 - 1, column=1,
            value="Investigation priority - a triage order, not a performance ranking").font = H2_FONT
    f2, l2 = write_table(ws, priority, row2, "tblPriority",
                         {"Active records": "#,##0", "Median age (days)": "#,##0",
                          "P90 age (days)": "#,##0", "Aged 90+": "#,##0",
                          "% of own queue aged 90+": "0.0", "% of city aged 90+": "0.0",
                          "Priority score": "0.0"})
    ws.conditional_formatting.add(
        f"I{f2 + 1}:I{l2}",
        ColorScaleRule(start_type="min", start_color="FFFFFFFF",
                       end_type="max", end_color="FFEB6834"))

    hotspots = agg("agg_priority_hotspots").head(40)[
        ["rank_by_aged_volume", "service_name", "council_district", "active_records",
         "aged_90_plus", "pct_aged_90_plus", "median_age_days", "aged_concentration_index"]
    ].rename(columns={
        "rank_by_aged_volume": "Rank", "service_name": "Service category",
        "council_district": "Council district", "active_records": "Active records",
        "aged_90_plus": "Aged 90+", "pct_aged_90_plus": "% aged 90+",
        "median_age_days": "Median age (days)",
        "aged_concentration_index": "Aged concentration index"})
    row3 = l2 + 3
    ws.cell(row=row3 - 1, column=1,
            value="Top 40 service x district cells by aged volume (cells with 200+ "
                  "active records)").font = H2_FONT
    write_table(ws, hotspots, row3, "tblHotspots",
                {"Active records": "#,##0", "Aged 90+": "#,##0", "% aged 90+": "0.0",
                 "Median age (days)": "#,##0", "Aged concentration index": "0.000"})

    # =====================================================================
    # 6. Data Quality
    # =====================================================================
    ws = wb.create_sheet("Data Quality")
    row = title_block(
        ws, "Data quality checks",
        "Generated by src/audit.py. A FAIL means the condition would invalidate a result "
        "if ignored - each one is worked around explicitly, see the Treatment column.",
        "Full report: 02_data_audit/data_quality_report.md")

    audit = json.loads(config.AUDIT_RESULTS.read_text())
    checks = pd.DataFrame([
        {"Check": f'{ch["check_id"]} {ch["name"]}',
         "Grade": ch["grade"],
         "Question": ch["question"],
         "Finding": ch["headline"],
         "Treatment": ch["treatment"]}
        for ch in audit["checks"]
    ])
    first, last = write_table(ws, checks, row, "tblAudit")
    for column in ("C", "D", "E"):
        ws.column_dimensions[column].width = 62
    for r in range(first + 1, last + 1):
        for column in ("C", "D", "E"):
            ws[f"{column}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 62

    for grade, color in (("PASS", "FFDDF3DD"), ("WARN", "FFFDF2D6"),
                         ("FAIL", "FFFBE0E0"), ("INFO", "FFE8EEF7")):
        ws.conditional_formatting.add(
            f"B{first + 1}:B{last}",
            CellIsRule(operator="equal", formula=[f'"{grade}"'],
                       fill=PatternFill("solid", fgColor=color)))

    summary_row = last + 2
    ws.cell(row=summary_row, column=1, value="Totals").font = H2_FONT
    for k, grade in enumerate(("PASS", "WARN", "FAIL", "INFO")):
        ws.cell(row=summary_row + 1 + k, column=1, value=grade)
        ws.cell(row=summary_row + 1 + k, column=2,
                value=f'=COUNTIF(tblAudit[Grade],"{grade}")').number_format = "0"

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"

    config.DASHBOARD_DIR.mkdir(parents=True, exist_ok=True)
    out = config.DASHBOARD_DIR / "san_diego_ops_review.xlsx"
    wb.save(out)
    return out


def main() -> int:
    out = build()
    print(f"  [out] {out.relative_to(config.REPO_ROOT)} "
          f"({out.stat().st_size / 1024:,.0f} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
